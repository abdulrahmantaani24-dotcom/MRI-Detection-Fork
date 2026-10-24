"""Read-only source audit; private reports/cache only. No downloads, GPU or training."""
from __future__ import annotations
import csv
import hashlib
import json
import os
import re
import sqlite3
import time
from collections import Counter, defaultdict
from pathlib import Path

DATASETS = {
    'normal': ('IXI/Healthy/IXI-T2', 'ixi'),
    'glioma': ('Glioma/Task01_BrainTumour', 'msd_task01'),
    'meningioma': ('DiseaseAcquisition/hwang_meningioma_5945963', 'hwang_meningioma_5945963'),
    'brain_metastasis': ('DiseaseAcquisition/proteas_23171182', 'proteas_23171182'),
    'post_stroke': ('DiseaseAcquisition/isles2022', 'isles2022'),
    'alzheimers': ('DiseaseAcquisition/oasis2_legacy', 'oasis2_legacy'),
}
STATUSES = {'READY', 'NEEDS_LABEL_REVIEW', 'NEEDS_MODALITY_REVIEW', 'NEEDS_FILE_VALIDATION', 'BLOCKED'}
POLICY_VERSION = 1


def save_json(path, value):
    """Only called on the verified private reports destination."""
    path = Path(path)
    if path.is_symlink():
        raise ValueError('Refusing report symlink')
    temporary = path.with_name(path.name + '.tmp')
    if temporary.is_symlink():
        raise ValueError('Refusing temporary report symlink')
    with temporary.open('w') as stream:
        json.dump(value, stream, indent=2)
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(path)


def load_json(path, errors, dataset):
    try:
        return json.loads(path.read_text())
    except Exception as exc:
        errors.append({'dataset': dataset, 'path': str(path), 'stage': 'metadata', 'error': str(exc)})
        return None


def is_mri(path):
    return path.name.lower().endswith(('.nii', '.nii.gz', '.img', '.img.gz'))


def metadata_first(root, key, errors):
    """Do this before walking large image trees. Never treat a template as reviewed."""
    metadata = {}
    for name in ['acquisition_report.json', 'source_and_license.json', 'file_manifest.json',
                 'patient_manifest.json', 'duplicate_image_report.json', 'dataset.json']:
        path = root / name
        if path.is_file() and not path.is_symlink():
            value = load_json(path, errors, key)
            expected_type = list if name == 'patient_manifest.json' else dict
            if value is not None and not isinstance(value, expected_type):
                errors.append({'dataset': key, 'path': str(path), 'stage': 'metadata', 'error': 'Unexpected metadata structure'})
                value = None
            metadata[name] = value
    template = root / 'patient_map_template.csv'
    if template.is_file() and not template.is_symlink():
        try:
            with template.open(newline='') as stream:
                metadata['patient_map_template.csv'] = list(csv.DictReader(stream))
        except Exception as exc:
            errors.append({'dataset': key, 'stage': 'metadata', 'error': str(exc), 'path': str(template)})
    # Clinical spreadsheets are inspected read-only, not interpreted as diagnoses.
    clinical = []
    for directory in [root, root / 'downloads']:
        if not directory.is_dir():
            continue
        for path in sorted(directory.glob('*.xlsx')):
            if path.is_symlink():
                continue
            try:
                import openpyxl
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
                try:
                    clinical.append({'path': str(path), 'sheets': [
                        {'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                         'header': [str(v) if v is not None else '' for v in next(sheet.iter_rows(values_only=True), ())]}
                        for sheet in book.worksheets], 'interpretation': 'Schema inspected only; clinical join and diagnostic adjudication require source review'})
                finally:
                    book.close()
            except Exception as exc:
                errors.append({'dataset': key, 'path': str(path), 'stage': 'clinical_metadata', 'error': str(exc)})
    metadata['clinical_spreadsheets'] = clinical
    return metadata


def read_reviewed_map(path, template, errors, key):
    if not path:
        return {}
    path = Path(path)
    if path.resolve() == template.resolve():
        errors.append({'dataset': key, 'stage': 'mapping', 'error': 'Template cannot be used as a reviewed mapping'})
        return {}
    allowed = {'meningioma': {'meningioma', 'brain_lesion_negative_lung_cancer'},
               'alzheimers': {'AD_clinical_baseline', 'control_stable', 'converter_excluded'}}
    result = {}
    try:
        with path.open(newline='') as stream:
            for row in csv.DictReader(stream):
                relative = row.get('relative_path', '')
                if not relative or relative.startswith('/') or '..' in Path(relative).parts or relative in result:
                    raise ValueError('Invalid or duplicate mapping path')
                if row.get('diagnosis') not in allowed.get(key, set()) or not all(row.get(k) for k in ['patient_id', 'visit_id', 'modality', 'role']):
                    raise ValueError(f'Incomplete source-reviewed mapping: {relative}')
                if row['role'] not in {'image', 'mask', 'ancillary'}:
                    raise ValueError(f'Unknown mapping role: {relative}')
                result[relative] = row
    except Exception as exc:
        errors.append({'dataset': key, 'stage': 'mapping', 'path': str(path), 'error': str(exc)})
        return {}
    return result


def identify(relative, key, mapping, shape):
    """Source case conventions only. Case IDs are not confirmed independent people."""
    name = Path(relative).name
    low = relative.lower()
    role, modalities, case, visit, diagnosis = 'image', ['unknown'], None, None, None
    if key == 'meningioma':
        role = 'unresolved'
    if key == 'normal':
        match = re.search(r'IXI\d+', name, re.I)
        case = match[0].upper() if match else None
        modalities = ['T2w'] if re.search(r'[-_]T2', name, re.I) else ['unknown']
        diagnosis = 'source_healthy_cohort_not_individually_confirmed'
    elif key == 'glioma':
        match = re.search(r'BRATS_\d+', name)
        case = match[0] if match else None
        role = 'mask' if 'labelsTr' in Path(relative).parts else 'image'
        modalities = ['mask'] if role == 'mask' else ['FLAIR', 'T1w', 'T1ce', 'T2w']
        diagnosis = 'unlabeled_test_not_training_positive' if 'imagesTs' in Path(relative).parts else 'glioma_cohort_pending_pair_validation'
    elif key == 'post_stroke':
        match = re.search(r'sub-strokecase\d+', relative)
        case = match[0] if match else None
        role = 'mask' if '_msk' in name else 'image'
        modalities = ['mask'] if role == 'mask' else [next((m for m in ['FLAIR', 'ADC', 'dwi'] if m.lower() in name.lower()), 'unknown')]
        diagnosis = 'acute_subacute_ischemic_stroke_not_chronic'
    elif key == 'brain_metastasis':
        match = re.search(r'(?:^|[/_])(P\d{2}[ab]?)(?:[/_.]|$)', relative, re.I)
        case = match[1][:3].upper() if match else None
        if case and not 1 <= int(case[1:]) <= 40:
            case = None
        vm = re.search(r'(?:^|[/_])(baseline|fu\d+)(?:[/_.]|$)', relative, re.I)
        course = match[1].upper() if match else None
        visit = f'{course}:{vm[1].lower()}' if vm else None
        if name.lower() in {'t1.nii.gz', 't1c.nii.gz', 't2.nii.gz', 'fla.nii.gz'} and 'brats' in low.split('/'):
            modalities = [{'t1.nii.gz': 'T1w', 't1c.nii.gz': 'T1ce', 't2.nii.gz': 'T2w', 'fla.nii.gz': 'FLAIR'}[name.lower()]]
        elif '_tumor_mask_' in low:
            role, modalities = 'mask', ['tumor_edema_necrosis']
        else:
            role, modalities = 'ancillary', ['CT_or_RTP_or_tissue_or_unresolved']
        diagnosis = 'brain_metastasis_source_cohort'
    elif key == 'alzheimers':
        match = re.search(r'OAS2_\d{4}', relative)
        case = match[0] if match else None
        vm = re.search(r'MR\d+', relative)
        visit = vm[0] if vm else None
        modalities = ['T1w']
    if key in {'alzheimers', 'meningioma'} and relative in mapping:
        row = mapping[relative]
        if case and row['patient_id'] not in {case, f'oasis2_legacy:{case}'}:
            raise ValueError('Reviewed mapping conflicts with OASIS source identity')
        case, diagnosis, visit, role, modalities = row['patient_id'], row['diagnosis'], row['visit_id'], row['role'], [row['modality']]
    return dict(case_id=case, visit_id=visit, diagnosis=diagnosis, role=role, modalities=modalities)


def fingerprint(path, paired_header=None):
    values = [str(path.resolve()), path.stat().st_size, path.stat().st_mtime_ns]
    if paired_header:
        values += [paired_header.stat().st_size, paired_header.stat().st_mtime_ns]
    return json.dumps(values)


def inspect_file(path, full_hash=False):
    import nibabel as nib
    import numpy as np
    image = nib.load(path)
    shape = list(image.shape)
    spacing = [float(z) for z in image.header.get_zooms()[:3]]
    if len(shape) < 3 or min(shape[:3]) < 2 or any(not 0 < z < 100 for z in spacing):
        raise ValueError('Invalid 3D shape or voxel spacing')
    if not np.isfinite(image.affine).all() or abs(np.linalg.det(image.affine[:3, :3])) < 1e-12:
        raise ValueError('Invalid/nonfinite/singular affine')
    result = {'shape': shape, 'spacing_mm': spacing, 'dtype': str(image.get_data_dtype()),
              'affine': image.affine.tolist(), 'sha256': None, 'paired_header_sha256': None}
    if full_hash:
        for source, field in [(path, 'sha256'), (Path(image.file_map['header'].filename) if 'header' in image.file_map else None, 'paired_header_sha256')]:
            if source:
                h = hashlib.sha256()
                with source.open('rb') as stream:
                    for chunk in iter(lambda: stream.read(1024**2), b''):
                        h.update(chunk)
                result[field] = h.hexdigest()
    return result


def msd_checks(root, rows, metadata, errors):
    description = metadata.get('dataset.json')
    def fail(message):
        errors.append({'dataset': 'glioma', 'stage': 'MSD', 'error': message})
    if not isinstance(description, dict):
        fail('Missing/invalid dataset.json')
        return
    declared = description.get('modality', {})
    aliases = {'flair': 'FLAIR', 't1w': 'T1w', 't1': 'T1w', 't1gd': 'T1ce', 't1ce': 'T1ce', 't2w': 'T2w', 't2': 'T2w'}
    if not isinstance(declared, dict) or [aliases.get(str(declared.get(str(i), '')).lower()) for i in range(4)] != ['FLAIR', 'T1w', 'T1ce', 'T2w']:
        fail('dataset.json channel order is missing/unexpected; do not assume modalities')
    counts = Counter(Path(r['relative_path']).parts[0] for r in rows)
    for folder, expected in [('imagesTr', 484), ('labelsTr', 484), ('imagesTs', 266)]:
        if counts[folder] != expected:
            fail(f'{folder}: expected {expected}, found {counts[folder]}')
    by_path = {r['relative_path']: r for r in rows}
    training = description.get('training', [])
    testing = description.get('test', [])
    if description.get('numTraining') != 484 or description.get('numTest') != 266 or len(training) != 484 or len(testing) != 266:
        fail('dataset.json counts/references disagree with the 484/266 public release')
    refs = set()
    for item in training:
        if not isinstance(item, dict):
            fail('Malformed training entry'); continue
        image, mask = item.get('image', '').removeprefix('./'), item.get('label', '').removeprefix('./')
        if image in refs or mask in refs:
            fail('Duplicate training reference')
        refs.update([image, mask])
        ir, mr = by_path.get(image), by_path.get(mask)
        if not ir or not mr or not image.startswith('imagesTr/') or not mask.startswith('labelsTr/'):
            fail(f'Missing/unsafe training image-label reference: {image}, {mask}')
            continue
        if ir['case_id'] != mr['case_id'] or ir['shape'][:3] != mr['shape'][:3] or ir['affine'] != mr['affine']:
            fail(f'Image-label identity/geometry mismatch: {image}')
    test_refs = []
    for item in testing:
        relative = item.removeprefix('./') if isinstance(item, str) else ''
        test_refs.append(relative)
        if relative not in by_path or not relative.startswith('imagesTs/'):
            fail(f'Missing/unsafe test reference: {relative}')
    if len(set(test_refs)) != len(test_refs) or set(test_refs) != {r['relative_path'] for r in rows if r['relative_path'].startswith('imagesTs/')}:
        fail('Test references are duplicated or incomplete')
    if refs != {r['relative_path'] for r in rows if r['relative_path'].startswith(('imagesTr/', 'labelsTr/'))}:
        fail('Training references do not cover training files exactly')
    for row in rows:
        if row['role'] == 'image' and (len(row['shape']) != 4 or row['shape'][3] != 4):
            fail(f'Expected four source channels: {row["relative_path"]}')


def link_duplicates(rows, metadata):
    """Conservative connected components, including retained acquisition split groups."""
    parents = {r['case_id']: r['case_id'] for r in rows if r['case_id']}
    def find(case):
        while parents[case] != case:
            parents[case] = parents[parents[case]]
            case = parents[case]
        return case
    def union(cases):
        cases = [c for c in cases if c in parents]
        for case in cases[1:]:
            parents[find(case)] = find(cases[0])
    by_hash, existing = defaultdict(list), defaultdict(list)
    for row in rows:
        if row.get('sha256') and row['role'] == 'image':
            by_hash[row['sha256']].append(row['case_id'])
        if row.get('existing_split_group'):
            existing[row['existing_split_group']].append(row['case_id'])
    for cases in list(by_hash.values()) + list(existing.values()):
        union(cases)
    # Known published-in-project ISLES finding: retain conservatively even without SHA scan.
    union(['sub-strokecase0090', 'sub-strokecase0169'])
    groups = defaultdict(list)
    for case in parents:
        groups[find(case)].append(case)
    for row in rows:
        component = sorted(groups[find(row['case_id'])]) if row['case_id'] else []
        row['split_group'] = component
        row['duplicate_link_review'] = len(component) > 1
    return [sorted(v) for v in groups.values() if len(v) > 1]


def audit(project_root, reports_root, *, reviewed_maps=None, full_hash=False, progress_every=25):
    """Call only in Colab for real data. Tests may use artificial temporary trees.

    Cache headers/hashes by path+size+mtime+paired header stats+mode+policy.
    It is an interruption aid, not an independent integrity proof. Errors aren't cached.
    """
    project, reports = Path(project_root).resolve(), Path(reports_root).resolve()
    sources = [(project / 'datasets' / relative).resolve() for relative, _ in DATASETS.values()]
    if any(reports == root or reports.is_relative_to(root) or root.is_relative_to(reports) for root in sources):
        raise ValueError('Reports/cache must be outside every dataset tree')
    if not reports.is_relative_to(project / 'reports'):
        raise ValueError('Use the private project reports directory')
    reports.mkdir(parents=True, exist_ok=True)
    cache_path = reports / 'readiness_cache.sqlite'
    if cache_path.is_symlink():
        raise ValueError('Refusing cache symlink')
    cache = sqlite3.connect(cache_path)
    if progress_every < 1:
        cache.close()
        raise ValueError('progress_every must be positive')
    cache.execute('CREATE TABLE IF NOT EXISTS headers (key TEXT PRIMARY KEY, value TEXT)')
    run = reports / ('run_' + time.strftime('%Y%m%dT%H%M%S') + '_' + str(time.time_ns()))
    run.mkdir()
    errors, summary, reviews, all_rows, compatibility = [], [], [], [], []
    metadata_by_key = {}
    # Metadata-first for ALL six datasets, before any expensive tree walk.
    for key, (relative, _) in DATASETS.items():
        root = project / 'datasets' / relative
        metadata_by_key[key] = metadata_first(root, key, errors)
    save_json(run / 'metadata_inventory.json', metadata_by_key)
    try:
        for key, (relative, dataset_id) in DATASETS.items():
            print(f'{key}: metadata read; auditing read-only', flush=True)
            root = project / 'datasets' / relative
            data = root / 'data' if (root / 'data').is_dir() else root
            metadata = metadata_by_key[key]
            before = len(errors)
            rows = []
            if not root.is_dir():
                errors.append({'dataset': key, 'stage': 'location', 'error': 'Dataset location missing', 'path': str(root)})
            mapping = read_reviewed_map((reviewed_maps or {}).get(key), root / 'patient_map_template.csv', errors, key)
            old_manifest = metadata.get('file_manifest.json') or {}
            if not isinstance(old_manifest, dict):
                old_manifest = {}
            old_rows = {r.get('relative_path'): r for r in old_manifest.get('files', []) if isinstance(r, dict) and isinstance(r.get('relative_path'), str)}
            if metadata.get('acquisition_report.json', {}) and metadata['acquisition_report.json'].get('status') == 'failed':
                errors.append({'dataset': key, 'stage': 'acquisition', 'error': 'Prior acquisition failed; review its private report before training'})
            seen = set()
            format_counts = Counter()
            if not root.resolve().is_relative_to((project / 'datasets').resolve()) or data.is_symlink():
                errors.append({'dataset': key, 'stage': 'location', 'error': 'Dataset tree redirects outside the expected source or data is a symlink'})
                root_available = False
            else:
                root_available = root.is_dir()
            for path in sorted(data.rglob('*')) if root_available else []:
                if path.is_file():
                    format_counts[''.join(path.suffixes).lower() or '[no extension]'] += 1
                if not is_mri(path) or not path.is_file() or any(p.startswith('.') or p == '__MACOSX' for p in path.relative_to(data).parts):
                    continue
                rel = path.relative_to(data).as_posix()
                seen.add(rel)
                try:
                    if path.is_symlink() or not path.resolve().is_relative_to(data.resolve()):
                        raise ValueError('Symlink/out-of-tree MRI refused')
                    paired = path.with_name(path.name.replace('.img', '.hdr')) if path.name.endswith(('.img', '.img.gz')) else None
                    if paired and paired.is_symlink():
                        raise ValueError('Paired header symlink refused')
                    token = str(POLICY_VERSION) + ':' + str(full_hash) + ':' + fingerprint(path, paired)
                    cached = cache.execute('SELECT value FROM headers WHERE key=?', (token,)).fetchone()
                    info = json.loads(cached[0]) if cached else inspect_file(path, full_hash)
                    if not cached:
                        cache.execute('INSERT OR REPLACE INTO headers VALUES (?,?)', (token, json.dumps(info)))
                        cache.commit()  # Each completed file survives interruption.
                    identity = identify(rel, key, mapping, info['shape'])
                    row = {'dataset': key, 'relative_path': rel, 'format': ''.join(path.suffixes),
                           'bytes': path.stat().st_size, **info, **identity}
                    row['existing_split_group'] = old_rows.get(rel, {}).get('split_group')
                    if row['sha256'] and old_rows.get(rel, {}).get('sha256') and row['sha256'] != old_rows[rel]['sha256']:
                        errors.append({'dataset': key, 'stage': 'integrity', 'path': rel, 'error': 'File SHA256 differs from prior acquisition manifest'})
                    rows.append(row)
                    reasons = []
                    if row['role'] == 'unresolved': reasons.append('image/mask role and modality unresolved')
                    if not row['case_id']: reasons.append('patient/case identity unresolved')
                    if row['diagnosis'] is None: reasons.append('diagnostic label unresolved')
                    if key in {'meningioma', 'alzheimers', 'brain_metastasis'} and not row['visit_id']: reasons.append('visit/course metadata unresolved')
                    if key == 'normal': reasons.append('source healthy cohort; individual clinical eligibility not confirmed')
                    if key == 'glioma' and row['diagnosis'] == 'unlabeled_test_not_training_positive': reasons.append('unlabeled test scan: exclude from confirmed-positive training')
                    if reasons:
                        reviews.append({'dataset': key, 'relative_path': rel, 'patient_or_case': row['case_id'], 'visit': row['visit_id'], 'diagnosis': row['diagnosis'], 'review_required': '; '.join(reasons)})
                except Exception as exc:
                    errors.append({'dataset': key, 'path': rel, 'stage': 'file', 'error': f'{type(exc).__name__}: {exc}'})
                if len(seen) % progress_every == 0:
                    print(f'{key}: {len(seen)} files visited; {len(errors)-before} new errors', flush=True)
            for missing in old_rows.keys() - seen:
                errors.append({'dataset': key, 'stage': 'manifest', 'path': missing, 'error': 'Manifest file missing from source tree'})
            for unused in mapping.keys() - seen:
                errors.append({'dataset': key, 'stage': 'mapping', 'path': unused, 'error': 'Reviewed map references missing MRI'})
            if not rows:
                errors.append({'dataset': key, 'stage': 'inventory', 'error': 'No readable MRI files; archives alone are not readiness'})
            if key == 'glioma':
                msd_checks(root, rows, metadata, errors)
            cases = {r['case_id'] for r in rows if r['case_id']}
            if key == 'post_stroke':
                if len(cases) != 250:
                    errors.append({'dataset': key, 'stage': 'ISLES', 'error': f'Expected 250 cases, found {len(cases)}'})
                for case in sorted(cases):
                    channels = Counter((r['role'], tuple(r['modalities'])) for r in rows if r['case_id'] == case)
                    expected = Counter({('image', ('FLAIR',)): 1, ('image', ('ADC',)): 1, ('image', ('dwi',)): 1, ('mask', ('mask',)): 1})
                    if channels != expected:
                        errors.append({'dataset': key, 'stage': 'ISLES', 'path': case, 'error': 'Expected exactly FLAIR, ADC, DWI and mask'})
            if key == 'brain_metastasis':
                if not metadata.get('clinical_spreadsheets'):
                    reviews.append({'dataset': key, 'relative_path': '[clinical metadata]', 'patient_or_case': '', 'visit': '', 'diagnosis': '', 'review_required': 'Clinical spreadsheet missing/unreadable; patient/course/visit and diagnosis join required'})
                else:
                    reviews.append({'dataset': key, 'relative_path': '[clinical metadata]', 'patient_or_case': '', 'visit': '', 'diagnosis': '', 'review_required': 'Clinical spreadsheet schema inventoried; verify each patient/course/visit join and source clinical inclusion before training'})
            if key == 'brain_metastasis' and len(cases) != 40:
                errors.append({'dataset': key, 'stage': 'PROTEAS', 'error': f'Expected 40 source patients, found {len(cases)}'})
            conflicts = defaultdict(set)
            for row in rows:
                if row['case_id'] and row['diagnosis']:
                    conflicts[row['case_id']].add(row['diagnosis'])
            for case, labels in conflicts.items():
                if len(labels) > 1:
                    errors.append({'dataset': key, 'stage': 'labels', 'path': case, 'error': f'Conflicting diagnoses: {sorted(labels)}'})
            template_rows = metadata.get('patient_map_template.csv', [])
            for template_row in template_rows:
                relative = template_row.get('relative_path', '')
                if relative not in mapping and relative not in {r['relative_path'] for r in rows}:
                    reviews.append({'dataset': key, 'relative_path': relative, 'patient_or_case': template_row.get('patient_id'), 'visit': template_row.get('visit_id'), 'diagnosis': template_row.get('diagnosis'), 'review_required': 'Unreviewed template entry; referenced MRI missing/unreadable or path does not match inventory; template values are not verified labels'})
            if key == 'alzheimers' and mapping:
                label_counts = Counter(next(iter(labels)) for labels in conflicts.values() if len(labels) == 1)
                if dict(label_counts) != {'AD_clinical_baseline': 64, 'control_stable': 72, 'converter_excluded': 14}:
                    errors.append({'dataset': key, 'stage': 'labels', 'error': 'Reviewed OASIS patient groups do not match published 64 AD / 72 stable controls / 14 converters'})
            duplicate_groups = link_duplicates(rows, metadata)
            for row in rows:
                if row['duplicate_link_review']:
                    reviews.append({'dataset': key, 'relative_path': row['relative_path'], 'patient_or_case': row['case_id'], 'visit': row['visit_id'], 'diagnosis': row['diagnosis'], 'review_required': 'duplicate-linked component; preserve identities; all channels/masks must stay together'})
            modalities = sorted({m for r in rows if r['role'] == 'image' for m in r['modalities']})
            dataset_errors = [e for e in errors if e['dataset'] == key]
            label_pending = any(r['dataset'] == key for r in reviews)
            status = 'BLOCKED' if not root_available or not rows else ('NEEDS_FILE_VALIDATION' if dataset_errors else ('NEEDS_LABEL_REVIEW' if label_pending else 'NEEDS_MODALITY_REVIEW'))
            acquisition = metadata.get('acquisition_report.json') or {}
            patients = defaultdict(set)
            for r in rows:
                if r['case_id'] and r['visit_id']:
                    patients[r['case_id']].add(r['visit_id'])
            summary.append({'dataset': key, 'dataset_id': dataset_id, 'status': status,
                'source_case_ids': len(cases), 'independent_clinically_confirmed_patients': 'not established',
                'MRI_scan_files': sum(r['role'] == 'image' for r in rows), 'masks': sum(r['role'] == 'mask' for r in rows),
                'ancillary_files': sum(r['role'] == 'ancillary' for r in rows),
                'unresolved_role_files': sum(r['role'] == 'unresolved' for r in rows),
                'image_cases_without_masks': sum(any(r['case_id'] == c and r['role'] == 'image' for r in rows) and not any(r['case_id'] == c and r['role'] == 'mask' for r in rows) for c in cases), 'modalities': ';'.join(modalities),
                'formats': ';'.join(sorted({r['format'] for r in rows})),
                'all_source_file_formats_counts': json.dumps(dict(format_counts)),
                'shape_counts': json.dumps(dict(Counter(str(r['shape']) for r in rows))),
                'spacing_mm_counts': json.dumps(dict(Counter(str(r['spacing_mm']) for r in rows))),
                'clinical_label_status': 'NEEDS_LABEL_REVIEW' if label_pending else 'source_cohort_only_pending_clinical_acceptance',
                'modality_status': 'NEEDS_MODALITY_REVIEW', 'payload_integrity_status': 'NEEDS_FILE_VALIDATION',
                'files_visited': len(seen),
                'errors': len(dataset_errors),
                'unreadable_or_invalid_files': sum(e['stage'] == 'file' for e in dataset_errors),
                'manifest_missing_files': sum(e['stage'] == 'manifest' for e in dataset_errors),
                'uncertain_diagnosis_files': sum(r['diagnosis'] is None for r in rows),
                'repeat_visit_cases': sum(len(v) > 1 for v in patients.values()),
                'duplicate_link_components': len(duplicate_groups), 'unresolved_identity_files': sum(not r['case_id'] for r in rows),
                'acquisition_training_ready': acquisition.get('training_ready', 'not recorded'),
                'IXI_claim_578_matches_actual': (len(rows) == 578) if key == 'normal' else '',
                'file_validation_scope': 'headers_and_full_file_SHA256' if full_hash else 'headers_only',
                'training_ready': False})
            compatibility.append({'dataset': key, 'modalities': ';'.join(modalities), 'reference': 'IXI T2w',
                'T2_channel_available': 'T2w' in modalities,
                'age_scanner_site_skullstrip_status': 'NOT_ADJUDICATED: published protocol/clinical join required; not inferred from shape',
                'binary_detector_compatibility': 'REVIEW_REQUIRED',
                'required_review': 'modality/channel selection; native vs registered space; scanner/site; age; resolution; skull stripping; cohort/source shortcuts',
                'clinical_scope': 'acute/subacute ischemia, not chronic post-stroke' if key == 'post_stroke' else ('clinical AD vs stable controls; exclude all converter visits; no MCI/dementia substitution' if key == 'alzheimers' else 'source diagnosis and inclusion criteria')})
            all_rows.extend(rows)
            save_json(run / f'{key}_file_inventory.json', rows)
            print(f'{key}: {status}; {len(rows)} readable files', flush=True)
        # Cross-dataset byte hashes are review links, never patient identity claims.
        hashes = defaultdict(list)
        for row in all_rows:
            if row['sha256'] and row['role'] == 'image':
                hashes[row['sha256']].append([row['dataset'], row['case_id'], row['relative_path']])
        cross_duplicates = [v for v in hashes.values() if len({r[0] for r in v}) > 1]
        save_json(run / 'cross_dataset_exact_duplicates.json', cross_duplicates)
        for items in cross_duplicates:
            for dataset, case, relative in items:
                reviews.append({'dataset': dataset, 'relative_path': relative, 'patient_or_case': case, 'visit': '', 'diagnosis': '', 'review_required': 'cross-dataset exact duplicate: manual review and one shared split component required'})
        write_reports(run, summary, reviews, compatibility, errors, full_hash)
        save_json(reports / 'latest_run.json', {'reports_directory': str(run), 'complete': True})
        return {'reports_directory': str(run), 'summary': summary, 'training_ready': False}
    finally:
        cache.close()


def write_csv(path, rows, fields):
    with path.open('x', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_reports(run, summary, reviews, compatibility, errors, full_hash):
    write_csv(run / 'Detection_Dataset_Summary.csv', summary, list(summary[0]))
    write_csv(run / 'Detection_Label_Review_Required.csv', reviews, ['dataset', 'relative_path', 'patient_or_case', 'visit', 'diagnosis', 'review_required'])
    write_csv(run / 'Detection_Modality_Compatibility.csv', compatibility, list(compatibility[0]))
    save_json(run / 'Detection_Validation_Errors.json', errors)
    save_json(run / 'Detection_Training_Readiness.json', {
        'training_ready': False, 'actual_audit_executed': True, 'statuses': {r['dataset']: r['status'] for r in summary},
        'allowed_statuses': sorted(STATUSES), 'full_file_hashing_enabled': full_hash,
        'limitations': ['Header reads do not prove voxel payload completeness or clinical validity.',
                        'Full-file SHA256 alone is not publisher archive verification or voxel-wise deduplication.',
                        'Cache uses size/mtime; same-size/same-mtime edits need a fresh cache.',
                        'No masks were loaded for value/lesion assessment; no preprocessing performed.',
                        'No modality/domain/clinical acceptance approval supplied; READY is not awarded automatically.'],
        'split_strategy': split_strategy()})
    lines = ['# Detection dataset readiness audit', '', 'No training or preprocessing performed. All original data were read-only.', '',
             '|Dataset|Status|Source case IDs|MRI files|Masks|Errors|', '|---|---|---:|---:|---:|---:|']
    lines += [f"|{r['dataset']}|{r['status']}|{r['source_case_ids']}|{r['MRI_scan_files']}|{r['masks']}|{r['errors']}|" for r in summary]
    lines += ['', 'Counts describe observed source files/case IDs, not independently confirmed healthy/diseased patients.',
              'See JSON inventories for per-file shape, spacing, affine, format, visit, label and conservative split components.',
              'Unlabeled MSD test images are excluded from confirmed-positive training eligibility. IXI count agreement does not establish individual health.',
              'Meningioma and OASIS templates are review queues, not verified mappings. Converters remain excluded from binary AD/control cohorts.',
              'PROTEAS source clinical/acquisition metadata must be reviewed; acquisition training_ready=false is never overridden.',
              'ISLES represents acute/subacute ischemic lesions rather than chronic post-stroke disease.', '',
              '## Compatibility and leakage',
              'IXI T2 controls versus T1-only OASIS and source-specific disease modalities can reveal the label via acquisition protocol. A mixed-source classifier is not justified by file availability. Review matched modality cohorts, age/site/scanner overlap, skull stripping, geometry/resolution and source balance; unknown protocol/clinical fields remain unresolved.',
              'Do not synthesize missing modalities or classify unknown diagnoses. Prefer source-balanced controls and external held-out sites. No cohort currently receives automatic READY.', '', '## Proposed splits (not assigned)', split_strategy(), '',
              '## Remaining file/clinical work',
              'Review errors and labels, source clinical metadata, publisher integrity evidence, payload/mask values and geometry, near duplicates and cross-source overlap. Hashing is optional and read-only, not a diagnosis test. Cache can resume interrupted scans; clear only the private cache to force re-reading. No source data or segmentation splits may be changed.']
    (run / 'Detection_Dataset_Readiness_Report.md').write_text('\n'.join(lines) + '\n')


def split_strategy():
    return ('After identity, diagnosis and modality review, build one global patient/source-overlap graph. '
            'Keep all repeated visits, treatment courses, channels, masks, slices and exact/near-duplicate-linked cases '
            'in one connected component; retain the ISLES 0090/0169 link without asserting shared identity. '
            'Respect any existing segmentation patient partitions; obtain an explicit cross-task overlap review before detection splitting. '
            'Exclude unresolved labels, OASIS converters and unlabeled MSD test cases. Prefer an external site-held-out test cohort; '
            'then consider 70/15/15 train/validation/test by verified component with disease/site/age/sex balance when documented. '
            'Small cohorts may need grouped cross-validation. Never tune on the held-out test cohort. '
            'No assignments are made by this audit, and no existing split/checkpoint is read as permission to change it.')
